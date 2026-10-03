"""Job Desk command line (Windows, macOS, Linux).

    python jobdesk/cli.py setup   --home DIR --owner "Firstname Lastname" [--port 8765] [--no-shortcut]
    python jobdesk/cli.py onboard --home DIR --cv FILE [--goals "..."]
    python jobdesk/cli.py start   --home DIR          # page + today's search
    python jobdesk/cli.py daily   --home DIR [--force]
    python jobdesk/cli.py tailor  --home DIR --folder YYYY-MM-DD/NN-slug
    python jobdesk/cli.py serve   --home DIR          # page only (foreground)
    python jobdesk/cli.py stats   --home DIR [--json] # usage and funnel (for the beta)
    python jobdesk/cli.py prep    --home DIR --folder YYYY-MM-DD/NN-slug   # interview prep pack
    python jobdesk/cli.py doctor  --home DIR [--quick]   # check the setup and say what to fix
    python jobdesk/cli.py export  --home DIR [--out FILE.xlsx]   # all picks and outcomes to Excel

--home can also come from the JOB_DESK_HOME environment variable.
"""
import argparse
import json
import os
import platform
import shutil
import socket
import subprocess
import sys
import time
import webbrowser
from datetime import date, datetime

CODE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
JD = os.path.join(CODE, "jobdesk")
# Shell allow-list: Python only (CV builder, fact check, JSON checks), through whichever shell Claude Code
# has: Bash when Git for Windows is installed, otherwise PowerShell. Measured behaviour in headless runs:
# network tools (curl) and other commands are refused; read-only commands and file writes inside the
# allowed folders still run (same reach as the file tools). Python itself can open network connections,
# so scan_audit() flags any shell command that looks network-capable. See docs/SECURITY.md.
PY_SHELL = "Bash(python:*),PowerShell(python *)"
SAFE_TOOLS = "Read,Write,Edit,Glob,Grep," + PY_SHELL + ",WebSearch,WebFetch"
BLOCKED_TOOLS = "Agent,Task,Workflow,Monitor"
SHELL_TOOLS = ("Bash", "PowerShell")


# ---------- helpers ----------

def home_dir(arg):
    h = arg or os.environ.get("JOB_DESK_HOME")
    if not h:
        sys.exit("Give --home or set JOB_DESK_HOME")
    return os.path.abspath(os.path.expanduser(h))


def load_config(home):
    try:
        with open(os.path.join(home, "config.json"), encoding="utf-8-sig") as f:
            return json.load(f)
    except FileNotFoundError:
        sys.exit(f"No config.json in {home}. Run: python jobdesk/cli.py setup --home \"{home}\" --owner \"Your Name\"")


def claude_exe():
    for c in (shutil.which("claude"), os.path.expanduser("~/.local/bin/claude.exe"), os.path.expanduser("~/.local/bin/claude")):
        if c and os.path.isfile(c):
            return c
    sys.exit("Claude Code not found. Install it from https://claude.com/claude-code and sign in once.")


def log_path(home, name):
    os.makedirs(os.path.join(home, "logs"), exist_ok=True)
    return os.path.join(home, "logs", name)


def run_agent(home, prompt, log, tools=SAFE_TOOLS, extra_dirs=(), blocked=BLOCKED_TOOLS):
    env = dict(os.environ, JOB_DESK_HOME=home, CLAUDE_CODE_PRINT_BG_WAIT_CEILING_MS="0",
               JOB_DESK_OWNER=load_config(home).get("owner", ""))
    # `python` in the agent's commands must be this Python (the installer's private copy has the CV libraries)
    env["PATH"] = os.path.dirname(sys.executable) + os.pathsep + env.get("PATH", "")
    cmd = [claude_exe(), "-p", prompt, "--allowedTools", tools, "--disallowedTools", blocked,
           "--permission-mode", "acceptEdits", "--output-format", "stream-json", "--verbose",
           "--add-dir", home, "--add-dir", CODE]
    for d in extra_dirs:
        cmd += ["--add-dir", d]
    audit = log[:-4] + ".audit.log" if log.endswith(".log") else log + ".audit"
    with open(log, "a", encoding="utf-8") as lf, open(audit, "a", encoding="utf-8") as af:
        stamp = datetime.now().isoformat(timespec="seconds")
        lf.write(f"=== started {stamp} ===" + chr(10)); af.write(f"=== started {stamp} ===" + chr(10))
        lf.flush()
        proc = subprocess.Popen(cmd, cwd=home, env=env, stdin=subprocess.DEVNULL, stdout=subprocess.PIPE,
                                stderr=subprocess.STDOUT, text=True, encoding="utf-8", errors="replace",
                                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        for line in proc.stdout:
            write_event(line, lf, af)
        code = proc.wait()
        stamp = datetime.now().isoformat(timespec="seconds")
        lf.write(f"=== finished {stamp}, exit {code} ===" + chr(10)); af.write(f"=== finished {stamp}, exit {code} ===" + chr(10))
    alerts = scan_audit(home, audit)
    if alerts:
        with open(log, "a", encoding="utf-8") as lf:
            lf.write(f"!!! {len(alerts)} security alert(s): see state/alerts.json and {os.path.basename(audit)}" + chr(10))
    return code


def sensitive_markers(home):
    """Strings that should never leave the machine inside a URL or search query."""
    import re
    markers = set()
    for name in ("cv_spec.json", "base_spec.json", "application_answers.md", "fact_bank.md"):
        path = os.path.join(home, "profile", name)
        if not os.path.isfile(path):
            continue
        text = open(path, encoding="utf-8-sig", errors="replace").read()
        markers.update(re.findall(r"[\w.+-]+@[\w-]+\.[\w.]+", text))                 # emails
        markers.update(m.replace(" ", "") for m in re.findall(r"\+\d[\d ]{7,}\d", text))  # phone numbers
        for line in text.splitlines():                                                  # distinctive fact lines
            words = re.findall(r"[A-Za-z]{4,}", line)
            if len(words) >= 8:
                markers.add(" ".join(words[:6]).lower())
    return {m for m in markers if len(m) >= 8}


def scan_audit(home, audit_path):
    """Flag web fetches or searches that carry personal data or suspiciously long queries."""
    from urllib.parse import unquote_plus
    markers = sensitive_markers(home)
    alerts = []
    try:
        lines = open(audit_path, encoding="utf-8", errors="replace").read().splitlines()
    except OSError:
        return alerts
    import re
    net = re.compile(r"https?://|urllib|requests\.|http\.client|socket|smtplib|ftplib|curl|wget|Invoke-WebRequest|Invoke-RestMethod", re.I)
    for line in lines:
        parts = line.split(" ", 2)
        if len(parts) >= 3 and parts[1].rstrip(":") in SHELL_TOOLS and net.search(parts[2]):
            alerts.append({"time": parts[0], "tool": parts[1].rstrip(":"), "target": parts[2][:300],
                           "reason": "shell command with network access"})
            continue
        if len(parts) < 3 or not parts[1].rstrip(":") in ("WebFetch", "WebSearch"):
            continue
        target = unquote_plus(parts[2])
        flat = target.replace(" ", "").lower()
        hit = next((m for m in markers if m.replace(" ", "").lower() in flat or m.lower() in target.lower()), None)
        long_query = parts[1].startswith("WebFetch") and "?" in target and len(target.split("?", 1)[1]) > 300
        if hit or long_query:
            alerts.append({"time": parts[0], "tool": parts[1].rstrip(":"), "target": target[:300],
                           "reason": "contains personal data" if hit else "unusually long query string"})
    if alerts:
        os.makedirs(os.path.join(home, "state"), exist_ok=True)
        path = os.path.join(home, "state", "alerts.json")
        try:
            existing = json.load(open(path, encoding="utf-8"))
        except (FileNotFoundError, json.JSONDecodeError):
            existing = []
        existing.append({"run": os.path.basename(audit_path), "at": datetime.now().isoformat(timespec="seconds"), "items": alerts})
        with open(path, "w", encoding="utf-8") as f:
            json.dump(existing, f, indent=2, ensure_ascii=False)
    return alerts


def write_event(line, lf, af):
    """Readable report goes to the log; every tool call (command, file, URL) goes to the audit log."""
    try:
        ev = json.loads(line)
    except json.JSONDecodeError:
        if line.strip():
            lf.write(line)
        return
    if ev.get("type") == "assistant":
        for part in (ev.get("message") or {}).get("content", []) or []:
            if part.get("type") == "tool_use":
                inp = part.get("input") or {}
                detail = (inp.get("command") or inp.get("url") or inp.get("query") or inp.get("file_path")
                          or inp.get("pattern") or json.dumps(inp)[:200])
                af.write(f"{datetime.now().strftime('%H:%M:%S')} {part.get('name')}: {str(detail)[:500]}" + chr(10))
        af.flush()
    elif ev.get("type") == "user":
        # Tool results: record refusals and errors so the audit shows what was blocked
        for part in (ev.get("message") or {}).get("content", []) or []:
            if isinstance(part, dict) and part.get("type") == "tool_result" and part.get("is_error"):
                body = part.get("content")
                if isinstance(body, list):
                    body = " ".join(x.get("text", "") for x in body if isinstance(x, dict))
                af.write(f"{datetime.now().strftime('%H:%M:%S')}   -> refused or failed: {str(body)[:160]}" + chr(10))
        af.flush()
    elif ev.get("type") == "result":
        lf.write((ev.get("result") or "") + chr(10))


def builder_and_spec(home, cfg):
    """The CV builder script and default CV spec file for this user (config can override both)."""
    builder = cfg.get("cv_builder") or os.path.join(JD, "cv", "build_cv.py")
    if not os.path.isabs(builder):
        builder = os.path.join(home, builder)
    return builder, cfg.get("cv_spec", "cv_spec.json")


def port_open(port):
    with socket.socket() as s:
        s.settimeout(0.5)
        return s.connect_ex(("127.0.0.1", port)) == 0


# ---------- commands ----------

def make_unix_launcher(home):
    """Double-clickable launcher: 'Job Desk.command' on macOS, a .desktop entry on Linux."""
    import shlex
    desktop = os.path.join(os.path.expanduser("~"), "Desktop")
    os.makedirs(desktop, exist_ok=True)
    cmd = f"{shlex.quote(sys.executable)} {shlex.quote(os.path.join(JD, 'cli.py'))} start --home {shlex.quote(home)}"
    if platform.system() == "Darwin":
        path = os.path.join(desktop, "Job Desk.command")
        with open(path, "w", encoding="utf-8") as f:
            f.write("#!/bin/sh\n" + cmd + "\n")
    else:
        path = os.path.join(desktop, "job-desk.desktop")
        with open(path, "w", encoding="utf-8") as f:
            f.write("[Desktop Entry]\nType=Application\nName=Job Desk\nTerminal=true\nExec=" + cmd + "\n")
    os.chmod(path, 0o755)
    return path


def cmd_setup(a):
    home = os.path.abspath(os.path.expanduser(a.home))
    if os.sep + ".claude" + os.sep in home + os.sep:
        sys.exit("Choose a data folder outside .claude (Claude Code protects it from automated edits).")
    for d in ("profile", "daily", "state", "logs"):
        os.makedirs(os.path.join(home, d), exist_ok=True)
    ex = os.path.join(CODE, "profile.example")
    for f in os.listdir(ex):
        dest = os.path.join(home, "profile", f)
        if not os.path.exists(dest):
            shutil.copy(os.path.join(ex, f), dest)
    cfg = os.path.join(home, "config.json")
    if not os.path.exists(cfg):
        with open(cfg, "w", encoding="utf-8") as f:
            json.dump({"owner": a.owner, "port": a.port, "networks": []}, f, indent=2)
    for rel, default in (("state/seen.json", {"jobs": {}}), ("state/decisions.json", {})):
        p = os.path.join(home, rel)
        if not os.path.exists(p):
            with open(p, "w", encoding="utf-8") as f:
                json.dump(default, f)
    subprocess.call([sys.executable, "-m", "pip", "install", "--quiet", "-r", os.path.join(CODE, "requirements.txt")])
    if not a.no_shortcut and platform.system() == "Windows":
        ps = shutil.which("pwsh") or shutil.which("powershell")
        desktop = os.path.join(os.path.expanduser("~"), "Desktop")
        lnk = os.path.join(desktop, "Job Desk.lnk")
        script = (f"$s=(New-Object -ComObject WScript.Shell).CreateShortcut('{lnk}');"
                  f"$s.TargetPath='{sys.executable}';$s.Arguments='\"{os.path.join(JD, 'cli.py')}\" start --home \"{home}\"';"
                  f"$s.WorkingDirectory='{home}';$s.Save()")
        subprocess.call([ps, "-NoProfile", "-Command", script])
    if not a.no_shortcut and platform.system() in ("Darwin", "Linux"):
        make_unix_launcher(home)
    print(f"Set up in {home}.\nNext: python \"{os.path.join(JD, 'cli.py')}\" onboard --home \"{home}\" --cv <your CV file> --goals \"what you want\"")


def cmd_onboard(a):
    home = home_dir(a.home)
    cfg = load_config(home)
    cv = os.path.abspath(os.path.expanduser(a.cv))
    prompt = (f"HOME is {home}. CODE is {CODE}. OWNER is {cfg.get('owner', '')}. CVFILE is {cv}. GOALS is: {a.goals}. "
              f"Follow the instructions in {os.path.join(CODE, 'prompts', 'onboard_prompt.md')} exactly.")
    print("Drafting your profile from your CV. This takes a few minutes.")
    code = run_agent(home, prompt, log_path(home, "onboard.log"), tools="Read,Write,Edit,Glob,Grep," + PY_SHELL,
                     extra_dirs=(os.path.dirname(cv),), blocked=BLOCKED_TOOLS + ",WebSearch,WebFetch")
    print(f"Done (exit {code}). Next: answer {os.path.join(home, 'profile', 'QUESTIONS.md')} and check profile/_preview.")
    return code


def cmd_daily(a):
    home = home_dir(a.home)
    cfg = load_config(home)
    today = date.today().isoformat()
    done = log_path(home, f"{today}.done")
    log = log_path(home, f"{today}.log")
    if os.path.exists(done) and not a.force:
        print("Today's search already ran."); return 0
    with open(log, "a", encoding="utf-8") as lf:  # flag open picks whose postings have closed
        subprocess.call([sys.executable, os.path.join(JD, "linkcheck.py"), "--home", home], stdout=lf, stderr=subprocess.STDOUT, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    if os.path.exists(os.path.join(home, "profile", "boards.json")):
        with open(log, "a", encoding="utf-8") as lf:
            subprocess.call([sys.executable, os.path.join(JD, "ingest.py"), "--home", home, "--date", today], stdout=lf, stderr=subprocess.STDOUT, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    # A second run on the same day writes to its own batch folder, so earlier picks are never overwritten
    batch = today
    for suffix in ("", "-b", "-c", "-d", "-e", "-f"):
        if not os.path.exists(os.path.join(home, "daily", today + suffix, "picks.json")):
            batch = today + suffix
            break
    prompt = (f"TODAY is {today}. BATCH is {batch}. HOME is {home}. CODE is {CODE}. OWNER is {cfg.get('owner', '')}. "
              f"Follow the instructions in {os.path.join(CODE, 'prompts', 'daily_prompt.md')} exactly, from step 0 to the end.")
    code = run_agent(home, prompt, log)
    if code == 0:
        open(done, "w").close()
    return code


def design_of(home, cfg):
    """The CV design the person chose in setup; the Job Desk design when none was chosen or a custom builder is set."""
    d = {"design": "jobdesk", "photo": "remove", "pages": 1, "improved": False}
    try:
        with open(os.path.join(home, "profile", "cv_design.json"), encoding="utf-8") as f:
            d.update(json.load(f))
    except (OSError, ValueError):
        pass
    if cfg.get("cv_builder"):
        d.update(design="jobdesk", photo="remove")
    return d


def build_jobdesk_preview(home, cfg, d):
    """Rebuild the master CV preview in the Job Desk design, with the photo if the person kept it."""
    spec_path = os.path.join(home, "profile", "cv_spec.json")
    try:
        with open(spec_path, encoding="utf-8") as f:
            spec = json.load(f)
    except (OSError, ValueError):
        print("No cv_spec.json yet."); return 1
    photo = None
    folder = os.path.join(home, "profile", "design")
    if d.get("photo") in ("keep", "auto") and os.path.isdir(folder):
        photo = next((n for n in sorted(os.listdir(folder)) if n.lower().startswith("photo.")), None)
    if photo:
        spec["photo"] = "design/" + photo
    else:
        spec.pop("photo", None)
    with open(spec_path, "w", encoding="utf-8") as f:
        json.dump(spec, f, indent=2, ensure_ascii=False)
    name = (cfg.get("owner") or spec.get("name") or "CV").replace(" ", "_") + "_CV"
    out = os.path.join(home, "profile", "_preview")
    code = subprocess.call([sys.executable, os.path.join(JD, "cv", "build_cv.py"), spec_path, out, "--name", name],
                           creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    return 0 if code in (0, 2) else code


def cmd_tailor(a):
    home = home_dir(a.home)
    cfg = load_config(home)
    import re
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}[A-Za-z0-9-]*/[A-Za-z0-9-]+", a.folder):
        sys.exit(f"Bad folder: {a.folder}")
    builder, spec = builder_and_spec(home, cfg)
    d = design_of(home, cfg)
    if d["design"] == "own" and not os.path.isfile(os.path.join(home, "profile", "design", "master.html")):
        d["design"] = "jobdesk"  # the own-design run hasn't finished: use the built-in design for now
    prompt = (f"HOME is {home}. CODE is {CODE}. OWNER is {cfg.get('owner', '')}. FOLDER is {a.folder}. "
              f"BUILDER is {builder}. BASESPEC is {spec}. DESIGN is {d['design']}. PHOTO is {d['photo']}. PAGES is {d['pages']}. "
              f"Follow the instructions in {os.path.join(CODE, 'prompts', 'tailor_prompt.md')} exactly.")
    code = run_agent(home, prompt, log_path(home, "tailor-" + a.folder.replace("/", "__") + ".log"),
                     tools="Read,Write,Edit,Glob,Grep," + PY_SHELL + ",WebFetch")
    ok = code == 0 and os.path.isfile(os.path.join(home, "daily", *a.folder.split("/"), "pack.json"))
    return 0 if ok else 1


def cmd_prep(a):
    """Interview prep pack for one job (called by the page server when a job is marked 'interview')."""
    home = home_dir(a.home)
    cfg = load_config(home)
    import re
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}[A-Za-z0-9-]*/[A-Za-z0-9-]+", a.folder):
        sys.exit(f"Bad folder: {a.folder}")
    _, spec = builder_and_spec(home, cfg)
    prompt = (f"HOME is {home}. CODE is {CODE}. OWNER is {cfg.get('owner', '')}. FOLDER is {a.folder}. BASESPEC is {spec}. "
              f"Follow the instructions in {os.path.join(CODE, 'prompts', 'interview_prompt.md')} exactly.")
    code = run_agent(home, prompt, log_path(home, "prep-" + a.folder.replace("/", "__") + ".log"))
    ok = code == 0 and os.path.isfile(os.path.join(home, "daily", *a.folder.split("/"), "interview.json"))
    return 0 if ok else 1


def cmd_doctor(a):
    """Check a setup end to end and say exactly what to fix."""
    import importlib.util
    import tempfile
    home = home_dir(a.home)
    results = []

    def check(name, ok, fix=""):
        results.append((ok, name, fix))

    check(f"Python {sys.version.split()[0]} (3.11+)", sys.version_info >= (3, 11), "Install Python 3.11 or newer")
    for mod, pkg in (("docx", "python-docx"), ("pymupdf", "pymupdf")):
        check(f"Python package {pkg}", importlib.util.find_spec(mod) is not None, f"python -m pip install {pkg}")
    claude = shutil.which("claude") or next((c for c in (os.path.expanduser("~/.local/bin/claude.exe"), os.path.expanduser("~/.local/bin/claude")) if os.path.isfile(c)), None)
    check("Claude Code installed", bool(claude), "Install from https://claude.com/claude-code, then run `claude` once to sign in")
    word = platform.system() == "Windows" and subprocess.call(
        [shutil.which("pwsh") or shutil.which("powershell") or "powershell", "-NoProfile", "-Command",
         "try { $w = New-Object -ComObject Word.Application; $w.Quit(); exit 0 } catch { exit 1 }"],
        stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL) == 0
    lo = bool(shutil.which("soffice") or shutil.which("libreoffice"))
    check("PDF exporter (Microsoft Word or LibreOffice)", word or lo, "Install LibreOffice (free) or Microsoft Word")
    cfg_path = os.path.join(home, "config.json")
    cfg = {}
    try:
        cfg = json.load(open(cfg_path, encoding="utf-8-sig"))
        check("config.json readable", True)
        check("config.json has your name", bool(cfg.get("owner")), 'Add "owner": "Firstname Lastname" to config.json')
    except (FileNotFoundError, json.JSONDecodeError) as e:
        check("config.json readable", False, f"Run setup, or fix {cfg_path} ({type(e).__name__})")
    builder, spec = builder_and_spec(home, cfg)
    for f in (spec, "fact_bank.md", "criteria.md", "application_answers.md"):
        path = os.path.join(home, "profile", f)
        exists = os.path.isfile(path)
        body = open(path, encoding="utf-8-sig", errors="replace").read() if exists else ""
        filled = exists and "\u2026" not in body[:4000] and "@example.com" not in body
        check(f"profile/{f} filled in", filled,
              "Run onboard, or replace the example text and the ... placeholders" if exists else "Run setup")
    if os.path.isfile(os.path.join(home, "profile", spec)) and (word or lo):
        out = tempfile.mkdtemp()
        r = subprocess.run([sys.executable, builder, os.path.join(home, "profile", spec), out, "--name", "doctor_check"],
                           capture_output=True, text=True, timeout=240)
        check("Default CV builds on one page", r.returncode == 0,
              "Shorten the summary or longest bullets in your default CV spec" if r.returncode == 2 else (r.stderr or r.stdout)[-200:])
    boards = os.path.join(home, "profile", "boards.json")
    if os.path.isfile(boards) and not a.quick:
        sys.path.insert(0, JD)
        import ingest
        cfgb = json.load(open(boards, encoding="utf-8-sig"))
        bad = []
        for b in cfgb.get("boards", [])[:40]:
            try:
                if b["ats"] == "workday":
                    ingest.fetch_workday(b, max_jobs=20)
                elif b["ats"] == "teamtailor":
                    ingest.fetch_teamtailor(ingest.FEEDS["teamtailor"].format(slug=b["slug"]))
                else:
                    ingest.fetch(ingest.FEEDS[b["ats"]].format(slug=b["slug"]))
            except Exception as e:  # noqa: BLE001
                bad.append(f"{b.get('company')} ({getattr(e, 'code', type(e).__name__)})")
        check(f"Company feeds respond ({len(cfgb.get('boards', [])) - len(bad)} ok)", not bad,
              "Fix or remove in profile/boards.json: " + ", ".join(bad[:8]))
    width = max(len(n) for _, n, _ in results) + 2
    for ok, name, fix in results:
        print(("  ok   " if ok else "  FIX  ") + name.ljust(width) + ("" if ok else "-> " + fix))
    failed = sum(1 for ok, _, _ in results if not ok)
    print(f"\n{'All good.' if not failed else str(failed) + ' thing(s) to fix.'}")
    return 1 if failed else 0


def cmd_export(a):
    """Write every pick with its status and outcome to an Excel workbook (needs openpyxl)."""
    import glob
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        sys.exit("Install openpyxl first: python -m pip install openpyxl")
    home = home_dir(a.home)
    decisions = {}
    try:
        decisions = json.load(open(os.path.join(home, "state", "decisions.json"), encoding="utf-8-sig"))
    except (FileNotFoundError, json.JSONDecodeError):
        pass
    closed = {}
    try:
        closed = json.load(open(os.path.join(home, "state", "closed.json"), encoding="utf-8-sig"))
    except (FileNotFoundError, json.JSONDecodeError):
        pass
    rows = []
    for pf in sorted(glob.glob(os.path.join(home, "daily", "*", "picks.json"))):
        try:
            data = json.load(open(pf, encoding="utf-8-sig"))
        except (json.JSONDecodeError, OSError):
            continue
        for j in data.get("jobs", []):
            d = decisions.get(j.get("id"), {})
            rows.append([j.get("date", ""), j.get("company", ""), j.get("role", ""), j.get("city", ""), j.get("track", ""),
                         j.get("tag", ""), j.get("score", ""), j.get("pay_estimate", ""), d.get("status", j.get("status", "new")),
                         d.get("skip_reason", ""), (d.get("submitted_at") or "")[:10], (d.get("outcome_at") or "")[:10],
                         "closed" if j.get("id") in closed else "", d.get("note", ""), j.get("url", "")])
    wb = Workbook()
    ws = wb.active
    ws.title = "Applications"
    head = ["Date", "Company", "Role", "City", "Track", "Tag", "Fit", "Pay estimate", "Status", "Skip reason",
            "Submitted", "Outcome date", "Posting", "Notes", "Link"]
    ws.append(head)
    for r in rows:
        ws.append(r)
        if r[-1].startswith("http"):
            c = ws.cell(row=ws.max_row, column=len(head))
            c.hyperlink = r[-1]
            c.font = Font(color="0563C1", underline="single")
    fill = PatternFill("solid", fgColor="1F3864")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
    for i, w in enumerate([11, 24, 44, 14, 11, 9, 6, 18, 11, 18, 11, 12, 9, 40, 50], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=c.column in (3, 14))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    st = wb.create_sheet("Funnel")
    from collections import Counter
    cnt = Counter(r[8] for r in rows)
    sent = sum(cnt[k] for k in ("submitted", "interview", "offer", "rejected"))
    for line in (("Picks shown", len(rows)), ("Approved (to submit)", cnt["apply"]), ("Applications sent", sent),
                 ("Interviews", cnt["interview"] + cnt["offer"]), ("Offers", cnt["offer"]), ("Rejected", cnt["rejected"]),
                 ("Skipped", cnt["skipped"]), ("Still to decide", cnt["new"])):
        st.append(line)
    st.column_dimensions["A"].width = 24
    out = os.path.abspath(a.out or os.path.join(home, f"Job Desk export {date.today().isoformat()}.xlsx"))
    wb.save(out)
    print(f"Wrote {len(rows)} picks to {out}")
    return 0


def cmd_apply(a):
    """Fold the user's answers into the profile, or apply a plain-language CV change request."""
    home = home_dir(a.home)
    cfg = load_config(home)
    if a.what not in ("answers", "revise", "goals", "design"):
        sys.exit("--what must be answers, revise, goals or design")
    d = design_of(home, cfg)
    if a.what == "design" and d["design"] == "jobdesk":
        return build_jobdesk_preview(home, cfg, d)  # no AI needed for the built-in design
    prompt = (f"HOME is {home}. CODE is {CODE}. OWNER is {cfg.get('owner', '')}. WHAT is {a.what}. "
              f"DESIGN is {d['design']}. PHOTO is {d['photo']}. PAGES is {d['pages']}. IMPROVED is {str(d.get('improved', False)).lower()}. "
              f"Follow the instructions in {os.path.join(CODE, 'prompts', 'apply_prompt.md')} exactly.")
    return run_agent(home, prompt, log_path(home, f"apply-{a.what}.log"), tools="Read,Write,Edit,Glob,Grep," + PY_SHELL,
                     blocked=BLOCKED_TOOLS + ",WebSearch,WebFetch")


def start_server(home, port):
    if port_open(port):
        return
    env = dict(os.environ, JOB_DESK_HOME=home, JOB_DESK_PORT=str(port))
    py = sys.executable
    if platform.system() == "Windows":
        pyw = os.path.join(os.path.dirname(py), "pythonw.exe")
        py = pyw if os.path.isfile(pyw) else py
        flags = subprocess.DETACHED_PROCESS | subprocess.CREATE_NEW_PROCESS_GROUP
        subprocess.Popen([py, os.path.join(JD, "server.py")], cwd=home, env=env, creationflags=flags, close_fds=True)
    else:
        subprocess.Popen([py, os.path.join(JD, "server.py")], cwd=home, env=env, start_new_session=True,
                         stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    for _ in range(20):
        if port_open(port):
            return
        time.sleep(0.25)


def cmd_start(a):
    home = home_dir(a.home)
    port = int(load_config(home).get("port", 8765))
    start_server(home, port)
    webbrowser.open(f"http://127.0.0.1:{port}")
    if os.path.exists(log_path(home, f"{date.today().isoformat()}.done")):
        print("Today's picks are already on the page."); time.sleep(3); return 0
    print("Finding today's jobs. This takes about 10 minutes; you can minimise this window.")
    a.force = False
    code = cmd_daily(a)
    print("Done. Refresh the page." if code == 0 else f"The run failed. See {log_path(home, date.today().isoformat() + '.log')}")
    if platform.system() == "Windows" and sys.stdin and sys.stdin.isatty():
        input("Press Enter to close")
    return code


def cmd_serve(a):
    home = home_dir(a.home)
    port = int(load_config(home).get("port", 8765))
    os.environ.update(JOB_DESK_HOME=home, JOB_DESK_PORT=str(port))
    sys.argv = [os.path.join(JD, "server.py")]
    import runpy
    runpy.run_path(os.path.join(JD, "server.py"), run_name="__main__")


def cmd_stats(a):
    """Usage and funnel for the beta: picks shown, decisions, applications, outcomes, per week."""
    import glob
    from collections import Counter
    home = home_dir(a.home)
    decisions = {}
    try:
        with open(os.path.join(home, "state", "decisions.json"), encoding="utf-8-sig") as f:
            decisions = json.load(f)
    except (FileNotFoundError, json.JSONDecodeError):
        pass
    jobs = []
    for pf in sorted(glob.glob(os.path.join(home, "daily", "*", "picks.json"))):
        try:
            with open(pf, encoding="utf-8-sig") as f:
                jobs += json.load(f).get("jobs", [])
        except (json.JSONDecodeError, OSError):
            continue
    for j in jobs:
        j.update(decisions.get(j.get("id"), {}))
    status = Counter(j.get("status", "new") for j in jobs)
    sent = sum(status[k] for k in ("submitted", "interview", "offer", "rejected"))
    decided = len(jobs) - status["new"]
    interviews = status["interview"] + status["offer"]
    days = sorted({j.get("date", "")[:10] for j in jobs if j.get("date")})
    skip_reasons = Counter()
    for j in jobs:
        if j.get("status") == "skipped":
            for r in (j.get("skip_reasons") or [j.get("skip_reason") or "no reason"]):
                skip_reasons[r] += 1
    by_track = Counter(j.get("track") or "Other" for j in jobs if j.get("status") in ("submitted", "interview", "offer", "rejected"))
    weeks = Counter()
    for j in jobs:
        if j.get("date"):
            try:
                y, w, _ = date.fromisoformat(j["date"][:10]).isocalendar()
                weeks[f"{y}-W{w:02d}"] += 1
            except ValueError:
                pass
    pct = lambda n, d: f"{round(100 * n / d)}%" if d else "-"
    report = {
        "active_days": len(days), "first_day": days[0] if days else None, "last_day": days[-1] if days else None,
        "picks_shown": len(jobs), "decided": decided, "decision_rate": pct(decided, len(jobs)),
        "approved": status["apply"] + sent, "approval_rate": pct(status["apply"] + sent, decided),
        "applications_sent": sent, "interviews": interviews, "interview_rate": pct(interviews, sent),
        "offers": status["offer"], "rejected": status["rejected"],
        "skip_reasons": dict(skip_reasons.most_common()), "applications_by_track": dict(by_track.most_common()),
        "picks_per_week": dict(sorted(weeks.items())),
    }
    if a.json:
        print(json.dumps(report, indent=2))
    else:
        print(f"Job Desk stats for {load_config(home).get('owner', home)}")
        print(f"  Active days: {report['active_days']} ({report['first_day']} to {report['last_day']})")
        print(f"  Picks shown: {report['picks_shown']}  |  decided {report['decided']} ({report['decision_rate']})")
        print(f"  Approved: {report['approved']} ({report['approval_rate']} of decided)")
        print(f"  Applications sent: {sent}  ->  interviews {interviews} ({report['interview_rate']})  ->  offers {status['offer']}")
        print(f"  Skip reasons: {report['skip_reasons'] or '-'}")
        print(f"  Applications by track: {report['applications_by_track'] or '-'}")
        print(f"  Picks per week: {report['picks_per_week'] or '-'}")
    return 0


def main():
    ap = argparse.ArgumentParser(description="Job Desk")
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("setup"); p.add_argument("--home", required=True); p.add_argument("--owner", required=True)
    p.add_argument("--port", type=int, default=8765); p.add_argument("--no-shortcut", action="store_true")
    p = sub.add_parser("onboard"); p.add_argument("--home"); p.add_argument("--cv", required=True); p.add_argument("--goals", default="")
    for name in ("start", "daily", "serve"):
        p = sub.add_parser(name); p.add_argument("--home"); p.add_argument("--force", action="store_true")
    p = sub.add_parser("tailor"); p.add_argument("--home"); p.add_argument("--folder", required=True)
    p = sub.add_parser("prep"); p.add_argument("--home"); p.add_argument("--folder", required=True)
    p = sub.add_parser("apply"); p.add_argument("--home"); p.add_argument("--what", required=True)
    p = sub.add_parser("doctor"); p.add_argument("--home"); p.add_argument("--quick", action="store_true")
    p = sub.add_parser("export"); p.add_argument("--home"); p.add_argument("--out")
    p = sub.add_parser("stats"); p.add_argument("--home"); p.add_argument("--json", action="store_true")
    a = ap.parse_args()
    sys.exit({"setup": cmd_setup, "onboard": cmd_onboard, "start": cmd_start, "daily": cmd_daily,
              "tailor": cmd_tailor, "serve": cmd_serve, "stats": cmd_stats, "prep": cmd_prep, "apply": cmd_apply, "doctor": cmd_doctor, "export": cmd_export}[a.cmd](a) or 0)


if __name__ == "__main__":
    main()
